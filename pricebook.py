"""Brendlangan narxlar ro'yxati — PDF va Excel fayl qilib chiqarish.

Ranglar: #e97609 (to'q sariq), #2e3239 (mokriy asfalt), #000000, #ffffff.
"""
import io
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdfcanvas

import cardmaker

ORANGE = colors.HexColor("#e97609")
ASPHALT = colors.HexColor("#2e3239")
BLACK = colors.HexColor("#000000")
WHITE = colors.HexColor("#ffffff")
ROW_ALT = colors.HexColor("#f4f5f6")
LINE = colors.HexColor("#dfe2e6")
GREY_TXT = colors.HexColor("#6b727c")

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "fonts")
_registered = False
CYRILLIC = __import__("re").compile(r"[\u0400-\u04FF]")


def _register_fonts() -> tuple[str, str]:
    """(bold, regular) shrift nomlarini qaytaradi."""
    global _registered
    bold_files = ["Montserrat-Bold.ttf", "Poppins-Bold.ttf", "DejaVuSans-Bold.ttf"]
    reg_files = ["Montserrat-Medium.ttf", "Montserrat-Regular.ttf",
                 "Poppins-Regular.ttf", "DejaVuSans.ttf"]
    if not _registered:
        for alias, files in (("BrandBold", bold_files), ("Brand", reg_files)):
            for f in files:
                path = os.path.join(FONT_DIR, f)
                if os.path.exists(path):
                    pdfmetrics.registerFont(TTFont(alias, path))
                    break
        # kirill matn uchun kafolatli shrift
        for alias, f in (("CyrBold", "DejaVuSans-Bold.ttf"), ("Cyr", "DejaVuSans.ttf")):
            path = os.path.join(FONT_DIR, f)
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(alias, path))
        _registered = True
    return "BrandBold", "Brand"


def _has_cyr(text: str) -> bool:
    return bool(CYRILLIC.search(text or ""))


def _money(v) -> str:
    s = str(v or "").strip().replace(" ", "").replace(",", "")
    if s.replace(".", "", 1).isdigit():
        return f"{int(float(s)):,}".replace(",", " ")
    return str(v or "").strip()


def group_by_category(products) -> "OrderedDict[str, list]":
    groups: "OrderedDict[str, list]" = OrderedDict()
    for p in products:
        cat = (p.get("category") or "Boshqa mahsulotlar").strip() or "Boshqa mahsulotlar"
        groups.setdefault(cat, []).append(p)
    for items in groups.values():
        items.sort(key=lambda x: (x.get("brand") or "", x.get("name") or ""))
    return OrderedDict(sorted(groups.items(), key=lambda kv: kv[0].lower()))


# ============================================================ PDF
class _Book:
    """A4 prays-list chizuvchi (reportlab canvas ustida qo'lda layout)."""

    M = 14 * mm            # chetki bo'shliq
    HEAD_H = 34 * mm       # sarlavha balandligi (1-sahifa)
    FOOT_H = 14 * mm
    ROW_H = 8.6 * mm
    CAT_H = 10 * mm

    def __init__(self, settings: dict):
        self.s = settings
        self.bold, self.reg = _register_fonts()
        self.buf = io.BytesIO()
        self.c = pdfcanvas.Canvas(self.buf, pagesize=A4)
        self.W, self.H = A4
        self.page = 0
        self.y = 0
        # ustun x koordinatalari
        left, right = self.M, self.W - self.M
        self.x_name = left + 4 * mm
        self.x_brand = right - 78 * mm
        self.x_unit = right - 46 * mm
        self.x_price = right - 4 * mm
        self.left, self.right = left, right

    # ---------------------------------------------------------- sahifa
    def _header(self, first: bool):
        c, s = self.c, self.s
        h = self.HEAD_H if first else 20 * mm
        c.setFillColor(ASPHALT)
        c.rect(0, self.H - h, self.W, h, stroke=0, fill=1)
        c.setFillColor(ORANGE)
        c.rect(0, self.H - h, self.W, 2.6 * mm, stroke=0, fill=1)

        logo = cardmaker.logo_path()
        has_logo = False
        tx = self.left
        if os.path.exists(logo):
            try:
                from PIL import Image as PILImage
                im = PILImage.open(logo)
                lh = 14 * mm if first else 9 * mm
                lw = lh * im.width / im.height
                c.drawImage(logo, self.left, self.H - h + (h - lh) / 2 + 1 * mm,
                            width=lw, height=lh, mask="auto")
                tx = self.left + lw + 6 * mm
                has_logo = True
            except Exception:
                pass
        else:
            box = 13 * mm if first else 9 * mm
            c.setFillColor(ORANGE)
            c.roundRect(self.left, self.H - h + (h - box) / 2 + 1 * mm, box, box, 2.5 * mm,
                        stroke=0, fill=1)
            c.setFillColor(WHITE)
            c.setFont(self.bold, box * 0.42)
            c.drawCentredString(self.left + box / 2, self.H - h + (h - box) / 2 + box * 0.33 + 1 * mm, "db")
            tx = self.left + box + 5 * mm

        # logo o'zida nom bo'lsa, yana takrorlamaymiz
        if not has_logo:
            c.setFillColor(WHITE)
            shop_name = s.get("shop_name") or "dunyabunya"
            c.setFont(self.f(shop_name, True), 17 if first else 12)
            c.drawString(tx, self.H - h + h * 0.58, shop_name)
            c.setFont(self.reg, 8.5 if first else 7.5)
            c.setFillColor(colors.HexColor("#b9bfc7"))
            c.drawString(tx, self.H - h + h * 0.58 - 5 * mm, "QURILISH MOLLARI")

        c.setFillColor(WHITE)
        c.setFont(self.bold, 12 if first else 9)
        c.drawRightString(self.right, self.H - h + h * 0.58, "NARXLAR RO'YXATI")
        c.setFont(self.reg, 8.5)
        c.setFillColor(ORANGE)
        c.drawRightString(self.right, self.H - h + h * 0.58 - 5 * mm, s.get("sana", ""))

        self.y = self.H - h - 9 * mm

    def _footer(self):
        c, s = self.c, self.s
        c.setStrokeColor(LINE)
        c.setLineWidth(0.6)
        c.line(self.left, self.FOOT_H + 3 * mm, self.right, self.FOOT_H + 3 * mm)
        c.setFont(self.reg, 8)
        c.setFillColor(GREY_TXT)
        c.drawString(self.left, self.FOOT_H - 2 * mm,
                     f"{s.get('shop_phone', '')}   ·   {s.get('channel_link', '')}")
        c.setFillColor(ASPHALT)
        c.setFont(self.bold, 8)
        c.drawRightString(self.right, self.FOOT_H - 2 * mm, f"{self.page}-bet")

    def new_page(self):
        if self.page:
            self._footer()
            self.c.showPage()
        self.page += 1
        self._header(first=self.page == 1)

    def _ensure(self, need: float):
        if self.y - need < self.FOOT_H + 8 * mm:
            self.new_page()

    # ---------------------------------------------------------- bloklar
    def table_head(self):
        c = self.c
        c.setFillColor(colors.HexColor("#eceef0"))
        c.rect(self.left, self.y - 6.6 * mm, self.right - self.left, 6.6 * mm, stroke=0, fill=1)
        c.setFillColor(GREY_TXT)
        c.setFont(self.bold, 7.6)
        base = self.y - 4.6 * mm
        c.drawString(self.x_name, base, "MAHSULOT")
        c.drawString(self.x_brand, base, "BREND")
        c.drawString(self.x_unit, base, "BIRLIK")
        c.drawRightString(self.x_price, base, "NARXI, SO'M")
        self.y -= 6.6 * mm

    def category(self, name: str, count: int):
        self._ensure(self.CAT_H + self.ROW_H * 2)
        c = self.c
        c.setFillColor(ORANGE)
        c.rect(self.left, self.y - self.CAT_H, self.right - self.left, self.CAT_H, stroke=0, fill=1)
        c.setFillColor(WHITE)
        c.setFont(self.f(name, True), 10.5)
        c.drawString(self.x_name, self.y - self.CAT_H + 3.4 * mm, name.upper())
        c.setFont(self.reg, 8.5)
        c.drawRightString(self.x_price, self.y - self.CAT_H + 3.4 * mm, f"{count} ta")
        self.y -= self.CAT_H
        self.table_head()

    def f(self, text: str, bold: bool = False) -> str:
        """Matnga mos shrift: kirill bo'lsa DejaVu, aks holda brend shrifti."""
        if _has_cyr(text):
            return "CyrBold" if bold else "Cyr"
        return self.bold if bold else self.reg

    def _fit(self, text: str, font: str, size: float, max_w: float) -> str:
        if self.c.stringWidth(text, font, size) <= max_w:
            return text
        while text and self.c.stringWidth(text + "…", font, size) > max_w:
            text = text[:-1]
        return text + "…"

    def row(self, p: dict, alt: bool):
        self._ensure(self.ROW_H)
        c = self.c
        top = self.y
        if alt:
            c.setFillColor(ROW_ALT)
            c.rect(self.left, top - self.ROW_H, self.right - self.left, self.ROW_H, stroke=0, fill=1)
        base = top - self.ROW_H + 2.9 * mm

        nm = p.get("name", "")
        fn = self.f(nm)
        c.setFillColor(BLACK)
        c.setFont(fn, 9.2)
        c.drawString(self.x_name, base,
                     self._fit(nm, fn, 9.2, self.x_brand - self.x_name - 4 * mm))

        br = p.get("brand", "") or "—"
        fb = self.f(br)
        c.setFillColor(GREY_TXT)
        c.setFont(fb, 8.4)
        c.drawString(self.x_brand, base,
                     self._fit(br, fb, 8.4, self.x_unit - self.x_brand - 3 * mm))
        un = (p.get("unit") or "—")[:10]
        c.setFont(self.f(un), 8.4)
        c.drawString(self.x_unit, base, un)

        old = _money(p.get("old_price"))
        price = _money(p.get("price")) or "—"
        c.setFillColor(ASPHALT)
        c.setFont(self.bold, 10)
        c.drawRightString(self.x_price, base, price)
        if old:
            w = c.stringWidth(price, self.bold, 10)
            c.setFont(self.reg, 7.4)
            c.setFillColor(GREY_TXT)
            ox = self.x_price - w - 4 * mm
            c.drawRightString(ox, base, old)
            ow = c.stringWidth(old, self.reg, 7.4)
            c.setStrokeColor(GREY_TXT)
            c.setLineWidth(0.5)
            c.line(ox - ow - 0.6 * mm, base + 0.9 * mm, ox + 0.6 * mm, base + 0.9 * mm)

        c.setStrokeColor(LINE)
        c.setLineWidth(0.4)
        c.line(self.left, top - self.ROW_H, self.right, top - self.ROW_H)
        self.y -= self.ROW_H

    def finish(self) -> bytes:
        self._footer()
        self.c.save()
        return self.buf.getvalue()


def make_pdf(products: list[dict], settings: dict) -> bytes:
    import db as _db

    s = dict(settings)
    s.setdefault("sana", f"{_db.now():%d.%m.%Y}")
    book = _Book(s)
    book.new_page()
    for cat, items in group_by_category(products).items():
        book.category(cat, len(items))
        for i, p in enumerate(items):
            book.row(p, alt=i % 2 == 1)
        book.y -= 4 * mm
    return book.finish()


# ============================================================ Excel
def make_xlsx(products: list[dict], settings: dict) -> bytes:
    import db as _db

    orange = "FFE97609"
    asphalt = "FF2E3239"
    alt = "FFF4F5F6"

    wb = Workbook()
    ws = wb.active
    ws.title = "Narxlar"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="FFDFE2E6")
    border = Border(bottom=thin)
    widths = [6, 54, 20, 12, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    shop = settings.get("shop_name") or "dunyabunya"
    ws.merge_cells("A1:F1")
    ws["A1"] = f"{shop} — NARXLAR RO'YXATI"
    ws["A1"].font = Font(name="Montserrat", bold=True, size=18, color="FFFFFFFF")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 46
    for col in "ABCDEF":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=asphalt)

    ws.merge_cells("A2:F2")
    ws["A2"] = (f"{_db.now():%d.%m.%Y}   ·   {settings.get('shop_phone', '')}"
                f"   ·   {settings.get('channel_link', '')}")
    ws["A2"].font = Font(name="Montserrat", bold=True, size=10, color="FFFFFFFF")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22
    for col in "ABCDEF":
        ws[f"{col}2"].fill = PatternFill("solid", fgColor=orange)

    logo = cardmaker.logo_path()
    if os.path.exists(logo):
        try:
            img = XLImage(logo)
            ratio = 52 / img.height
            img.height, img.width = 52, int(img.width * ratio)
            ws.add_image(img, "E1")
        except Exception:
            pass

    headers = ["№", "Mahsulot nomi", "Brend", "Birlik", "Eski narx", "Narxi, so'm"]
    r = 4
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=i, value=h)
        cell.font = Font(name="Montserrat", bold=True, size=10, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=asphalt)
        cell.alignment = Alignment(horizontal="center" if i != 2 else "left", vertical="center")
    ws.row_dimensions[r].height = 24
    ws.freeze_panes = f"A{r + 1}"

    r += 1
    n = 0
    for cat, items in group_by_category(products).items():
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value=f"{cat.upper()}   ({len(items)} ta)")
        c.font = Font(name="Montserrat", bold=True, size=11, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=orange)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 22
        r += 1

        for i, p in enumerate(items):
            n += 1
            fill = PatternFill("solid", fgColor=alt) if i % 2 else None
            vals = [n, p.get("name", ""), p.get("brand", "") or "—", p.get("unit", "") or "—"]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=col, value=v)
                cell.font = Font(name="Montserrat", size=10,
                                 color="FF6B727C" if col in (1, 3, 4) else "FF000000")
                cell.alignment = Alignment(
                    horizontal="left" if col == 2 else "center", vertical="center")
                cell.border = border
                if fill:
                    cell.fill = fill

            for col, key, bold in ((5, "old_price", False), (6, "price", True)):
                raw = str(p.get(key) or "").replace(" ", "")
                cell = ws.cell(row=r, column=col,
                               value=int(float(raw)) if raw.replace(".", "", 1).isdigit() else (raw or None))
                cell.number_format = "#,##0"
                cell.font = Font(name="Montserrat", size=11 if bold else 9, bold=bold,
                                 color="FF2E3239" if bold else "FF9AA1AA",
                                 strike=not bold and bool(raw))
                cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
                cell.border = border
                if fill:
                    cell.fill = fill
            ws.row_dimensions[r].height = 20
            r += 1

    ws.auto_filter.ref = f"A4:F{r - 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
