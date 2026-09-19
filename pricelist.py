"""Excel / CSV narxlar ro'yxatini o'qish.

Ustun nomlari o'zbekcha, ruscha yoki inglizcha bo'lishi mumkin — bot o'zi taniydi.
Kerakli ustun faqat bittasi: mahsulot nomi. Qolganlari ixtiyoriy.
"""
import csv
import io
import re

from openpyxl import load_workbook

COLUMN_ALIASES = {
    "name": [
        "nom", "nomi", "mahsulot", "mahsulot nomi", "tovar", "tovar nomi", "maxsulot",
        "название", "наименование", "товар", "продукт", "name", "product", "title",
        "отображаемое имя", "display name", "полное наименование",
    ],
    "price": [
        "narx", "narxi", "narh", "yangi narx", "sotuv narxi", "summa",
        "chakana", "chakana narx", "chakana narxi", "dona narxi",
        "цена", "стоимость", "розничная", "розница", "розничная цена",
        "price", "new price", "cost", "retail",
    ],
    "wholesale": [
        "ulgurji", "ulgurji narx", "ulgurji narxi", "optom",
        "оптовая", "опт", "оптовая цена", "wholesale",
    ],
    "old_price": [
        "eski narx", "eski narxi", "avvalgi narx", "старая цена", "old price", "было",
    ],
    "unit": [
        "birlik", "birligi", "olcham", "o'lchov", "olchov", "olchov birligi",
        "ед", "ед.изм", "ед. изм.", "единица", "unit", "uom",
    ],
    "category": [
        "kategoriya", "turkum", "guruh", "bolim", "bo'lim",
        "категория", "группа", "раздел", "category", "group",
    ],
    "brand": [
        "brend", "brand", "firma", "ishlab chiqaruvchi", "zavod",
        "бренд", "производитель", "марка", "manufacturer",
    ],
    "note": [
        "izoh", "tavsif", "qoshimcha", "описание", "комментарий", "примечание",
        "note", "description", "comment",
    ],
}


def _norm_header(h) -> str:
    s = str(h or "").strip().lower()
    s = s.replace("ʻ", "'").replace("ʼ", "'").replace("`", "'").replace("’", "'")
    s = re.sub(r"[^\w\s.']+", " ", s, flags=re.UNICODE)
    return re.sub(r"\s+", " ", s).strip()


def _map_columns(header_row) -> dict:
    mapping = {}
    for idx, cell in enumerate(header_row):
        h = _norm_header(cell)
        if not h:
            continue
        for field, aliases in COLUMN_ALIASES.items():
            if field in mapping:
                continue
            if h in aliases or any(h.startswith(a) for a in aliases):
                mapping[field] = idx
                break
    return mapping


def _find_header(rows) -> tuple[int, dict]:
    """Sarlavha qatorini birinchi 10 qator ichidan topadi."""
    best_i, best_map = -1, {}
    for i, row in enumerate(rows[:10]):
        m = _map_columns(row)
        if "name" in m and len(m) > len(best_map):
            best_i, best_map = i, m
    return best_i, best_map


CODE_PREFIX = re.compile(r"^\s*[\[\(]\s*[A-Za-z0-9][A-Za-z0-9._/-]*\s*[\]\)]\s*")
UPPER_BRAND = re.compile(r"\b([A-ZА-ЯЎҚҒҲ]{3,}(?:[- ][A-ZА-ЯЎҚҒҲ]{2,})?)\b")
UNIT_LIKE = {"MM", "SM", "KG", "GR", "ML", "LED", "PVC", "UV", "GKL", "GKLV", "OSB", "MDF"}


def clean_name(raw: str) -> str:
    """'[0066-01890] Bazalt PETRAWOOL (100mm)' -> 'Bazalt PETRAWOOL (100mm)'"""
    s = str(raw or "").strip()
    prev = None
    while s != prev:                 # bir nechta kod ketma-ket bo'lishi mumkin
        prev = s
        s = CODE_PREFIX.sub("", s)
    return re.sub(r"\s{2,}", " ", s).strip(" -–—|")


def guess_category(name: str) -> str:
    """Kategoriya ustuni bo'lmasa — nomning birinchi so'zidan."""
    first = clean_name(name).split()
    if not first:
        return ""
    w = first[0].strip(".,;:")
    return w.capitalize() if len(w) > 2 else ""


def guess_brand(name: str, category: str = "") -> str:
    """Brend ustuni bo'lmasa — nomdagi BOSH HARFLI so'z (PETRAWOOL, EVEREST)."""
    body = clean_name(name)
    if category:
        body = re.sub(r"^" + re.escape(category), "", body, flags=re.IGNORECASE).strip()
    for m in UPPER_BRAND.finditer(body):
        cand = m.group(1).strip()
        if cand.upper() in UNIT_LIKE or cand.isdigit():
            continue
        return cand.title() if len(cand) > 3 else cand
    return ""


def _clean_price(v) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, (int, float)):
        return str(int(round(float(v))))
    s = str(v).strip()
    s = re.sub(r"(so'm|som|сум|sum|uzs|руб)\.?", "", s, flags=re.IGNORECASE).strip()
    digits = re.sub(r"[^\d]", "", s.replace(".", "").replace(",", ""))
    return digits or s


def _rows_from_xlsx(data: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        sheet_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if len(sheet_rows) > len(rows):
            rows = sheet_rows
    wb.close()
    return rows


def _rows_from_csv(data: bytes) -> list[list]:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def parse(data: bytes, filename: str, default_category: str = "") -> tuple[list[dict], str]:
    """Faylni o'qib, mahsulotlar ro'yxatini qaytaradi. (mahsulotlar, xabar)

    default_category — fayl bilan birga yozilgan kategoriya nomi (caption).
    Bo'sh bo'lsa, kategoriya mahsulot nomining birinchi so'zidan olinadi.
    """
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xltx")):
        rows = _rows_from_xlsx(data)
    elif lower.endswith((".csv", ".txt", ".tsv")):
        rows = _rows_from_csv(data)
    else:
        return [], (
            "❌ Bu format qo'llab-quvvatlanmaydi.\n"
            "Iltimos <b>.xlsx</b> yoki <b>.csv</b> fayl yuboring.\n"
            "(PDF bo'lsa — Excel'ga o'girib yuboring.)"
        )

    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        return [], "❌ Fayl bo'sh ko'rinadi."

    h_idx, mapping = _find_header(rows)
    if h_idx < 0:
        return [], (
            "❌ Mahsulot nomi ustunini topa olmadim.\n\n"
            "Birinchi qatorda ustun nomlari bo'lsin, masalan:\n"
            "<code>Nomi | Kategoriya | Brend | Narx | Birlik | Eski narx | Izoh</code>"
        )

    items, skipped = [], 0
    for row in rows[h_idx + 1:]:
        def cell(field):
            i = mapping.get(field)
            if i is None or i >= len(row):
                return ""
            v = row[i]
            return "" if v is None else str(v).strip()

        name = clean_name(cell("name"))
        if not name or _norm_header(name) in COLUMN_ALIASES["name"]:
            continue
        price = _clean_price(row[mapping["price"]]) if "price" in mapping and mapping["price"] < len(row) else ""
        if not price and "wholesale" in mapping and mapping["wholesale"] < len(row):
            price = _clean_price(row[mapping["wholesale"]])
        old = _clean_price(row[mapping["old_price"]]) if "old_price" in mapping and mapping["old_price"] < len(row) else ""
        if not price:
            skipped += 1
        category = cell("category") or default_category or guess_category(name)
        brand = cell("brand") or guess_brand(name, category)
        items.append({
            "name": name,
            "price": price,
            "old_price": old,
            "unit": cell("unit"),
            "category": category,
            "brand": brand,
            "note": cell("note"),
        })

    found = ", ".join(sorted(mapping)) or "—"
    msg = f"Topilgan ustunlar: <code>{found}</code>"
    if skipped:
        msg += f"\n⚠️ {skipped} ta qatorda narx yo'q — ular navbatga chiqmaydi."
    return items, msg
